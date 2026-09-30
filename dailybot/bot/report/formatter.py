"""
report/formatter.py — .docx formatter for the weekly report (Slice 5).

Ported near-verbatim from Teleport's bot/report/formatter.py. Only OUTPUT_DIR
and the HARIAN_PROJECT import changed: output now lands in dailybot's
data/reports/ and the constant is inlined (v2 has no database module).
"""

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

from bot import config
from bot.report.utils import BM_DAYS, BM_MONTHS

HARIAN_PROJECT = "Harian"


def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _section_heading(doc: Document, number: str, title: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{number}. {title.upper()}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def _add_bullet_lines(doc: Document, text: str) -> None:
    for line in text.strip().splitlines():
        line = line.strip()
        if not line:
            continue
        clean = re.sub(r"^[\*\-\•]\s*", "", line)
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(clean).font.size = Pt(11)


def _add_plain_text(doc: Document, text: str) -> None:
    for block in text.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        has_bullets = all(re.match(r"^[\*\-\•]", l.strip()) for l in lines if l.strip())
        if has_bullets:
            _add_bullet_lines(doc, block)
        else:
            p = doc.add_paragraph(block)
            if p.runs:
                p.runs[0].font.size = Pt(11)


def build_docx(
    report_sections: dict[str, str],
    week_line: str,
    date_range_line: str,
    activities_by_day: dict[str, list[dict]],
    working_days: list[date],
) -> Path:
    """Build the full report .docx and return its path."""
    output_dir = config.REPORTS_DIR  # read at call time so tests/sims can redirect it
    output_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("LAPORAN MINGGUAN PEKERJA")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{week_line}\n{date_range_line}").font.size = Pt(12)

    doc.add_paragraph()

    info_table = doc.add_table(rows=2, cols=2)
    info_table.style = "Table Grid"
    labels = [("Nama", "Muhammad Naqib"), ("Jawatan", "Pekerja Pengeluaran")]
    for i, (label, value) in enumerate(labels):
        row = info_table.rows[i]
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = value
    doc.add_paragraph()

    _section_heading(doc, "1", "Ringkasan Kerja Minggu Ini")
    _add_plain_text(doc, report_sections.get("ringkasan", "—"))
    doc.add_paragraph()

    _section_heading(doc, "2", "Isu / Masalah Dihadapi")
    _add_plain_text(doc, report_sections.get("isu", "—"))
    doc.add_paragraph()

    _section_heading(doc, "3", "Cadangan / Keperluan Sokongan")
    _add_plain_text(doc, report_sections.get("cadangan", "—"))
    doc.add_paragraph()

    _section_heading(doc, "4", "Pencapaian")

    p = doc.add_paragraph()
    p.add_run("Pencapaian dalam kerja:").bold = True
    _add_plain_text(doc, report_sections.get("pencapaian_kerja", "—"))

    p2 = doc.add_paragraph()
    p2.add_run("Pencapaian peribadi:").bold = True
    _add_plain_text(doc, report_sections.get("pencapaian_peribadi", "—"))
    doc.add_paragraph()

    _section_heading(doc, "5", "Refleksi")

    p3 = doc.add_paragraph()
    p3.add_run("Apa yang saya belajar / sedar minggu ini:").bold = True
    _add_plain_text(doc, report_sections.get("refleksi_belajar", "—"))

    p4 = doc.add_paragraph()
    p4.add_run("Apa saya nak perbaiki minggu depan:").bold = True
    _add_plain_text(doc, report_sections.get("refleksi_perbaiki", "—"))
    doc.add_paragraph()

    _section_heading(doc, "6", "Laporan Aktiviti")
    doc.add_paragraph()

    act_table = doc.add_table(rows=1, cols=3)
    act_table.style = "Table Grid"
    act_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Cm(4), Cm(10.5), Cm(3)]
    for i, width in enumerate(widths):
        for row in act_table.rows:
            row.cells[i].width = width

    hdr = act_table.rows[0]
    headers = ["Hari / Tarikh", "Aktiviti", "Waktu Balik"]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        _set_cell_bg(cell, "1F4E79")
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(11)

    for d in working_days:
        key = d.isoformat()
        acts = activities_by_day.get(key, [])
        day_bm = BM_DAYS[d.weekday()]
        month_bm = BM_MONTHS[d.month]
        date_str = f"{day_bm}\n{d.day} {month_bm} {d.year}"

        row = act_table.add_row()
        row.cells[0].text = date_str
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(11)

        act_cell = row.cells[1]
        if not acts:
            p = act_cell.paragraphs[0]
            p.add_run("Tiada aktiviti direkodkan.").italic = True
            p.runs[0].font.size = Pt(11)
        else:
            first = True
            for a in acts:
                _raw = a.get("detail")
                if a.get("project") == HARIAN_PROJECT and str(_raw or "").strip() in ("", "-"):
                    detail_part = " — Tiada"
                else:
                    detail_part = f" — {_raw}" if _raw else ""
                sentence = f"{a['project']}: {a['activity_type']}{detail_part}"
                if first:
                    p = act_cell.paragraphs[0]
                    p.text = ""
                    first = False
                else:
                    p = act_cell.add_paragraph()
                run = p.add_run(f"• {sentence}")
                run.font.size = Pt(11)

        row.cells[2].text = ""

    safe_range = date_range_line.replace(" ", "_").replace("-", "-")
    filename = f"Laporan_{safe_range}.docx"
    out_path = output_dir / filename
    doc.save(str(out_path))
    return out_path